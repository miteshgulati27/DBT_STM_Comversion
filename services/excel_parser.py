import time
from pathlib import Path
from openpyxl import load_workbook

KNOWN_HEADERS = {
    "target_column": ["target column", "target field", "target column name", "column name", "target_column", "stm column"],
    "source_table": ["source table", "source table name", "source_table", "source entity"],
    "source_column": ["source column", "source field", "source column name", "source_column", "column mapping", "column_mapping", "source column mapping"],
    "data_type": ["data type", "datatype", "data_type", "type", "target data type", "target datatype"],
    "transformation": ["transformation", "transformation logic", "transform", "business logic", "logic", "from/join/where", "from / join / where"],
    "business_rule": ["business rule", "rule", "rule name", "business_rule", "general rule applied", "general rule", "cleansing rule"],
    "scd_type": ["scd type", "scd_type", "scd", "scd type"],
}


SKIP_TABS = {
    "domains", "table of contents", "instructions", "versions",
    "template", "sample", "business rules", "conformed",
    "reference", "audit", "coverable type master list",
    "sheet1", "backward compatibility", "issues list",
    "progress report", "cdc",
}

_stm_cache = {}
CACHE_TTL = 300  # 5 minutes


def parse_stm_workbook(stm_path: Path) -> dict:
    cache_key = str(stm_path)
    now = time.time()

    if cache_key in _stm_cache and (now - _stm_cache[cache_key]["time"]) < CACHE_TTL:
        return _stm_cache[cache_key]["tabs"]

    wb = load_workbook(str(stm_path), read_only=True, data_only=True)
    tabs = {}
    for sheet_name in wb.sheetnames:
        if sheet_name.lower().strip() in SKIP_TABS:
            continue
        ws = wb[sheet_name]
        headers = _detect_headers(ws)
        if headers and headers.get("target_column") is not None:
            row_count = 0
            for row in ws.iter_rows(min_row=headers["_header_row"] + 1, values_only=True):
                if row[headers["target_column"]]:
                    row_count += 1
            tabs[sheet_name] = {
                "column_count": row_count,
                "headers": {k: v for k, v in headers.items() if not k.startswith("_")},
                "header_row": headers["_header_row"]
            }
    wb.close()

    _stm_cache[cache_key] = {"time": now, "tabs": tabs}
    return tabs


def get_tab_columns(stm_path: Path, tab_name: str) -> list:
    cache_key = f"{stm_path}::{tab_name}"
    now = time.time()

    if cache_key in _stm_cache and (now - _stm_cache[cache_key]["time"]) < CACHE_TTL:
        return _stm_cache[cache_key]["columns"]

    wb = load_workbook(str(stm_path), read_only=True, data_only=True)
    if tab_name not in wb.sheetnames:
        wb.close()
        return []

    ws = wb[tab_name]
    headers = _detect_headers(ws)
    if not headers or headers.get("target_column") is None:
        wb.close()
        return []

    columns = []
    for row in ws.iter_rows(min_row=headers["_header_row"] + 1, values_only=True):
        target_col = row[headers["target_column"]] if headers.get("target_column") is not None else None
        if not target_col:
            continue

        col_data = {
            "target_column": str(target_col).strip() if target_col else "",
            "source_table": str(row[headers["source_table"]]).strip() if headers.get("source_table") is not None and row[headers["source_table"]] else "",
            "source_column": str(row[headers["source_column"]]).strip() if headers.get("source_column") is not None and row[headers["source_column"]] else "",
            "data_type": str(row[headers["data_type"]]).strip() if headers.get("data_type") is not None and row[headers["data_type"]] else "",
            "transformation": str(row[headers["transformation"]]).strip() if headers.get("transformation") is not None and row[headers["transformation"]] else "",
            "business_rule": str(row[headers["business_rule"]]).strip() if headers.get("business_rule") is not None and row[headers["business_rule"]] else "",
            "scd_type": str(row[headers["scd_type"]]).strip() if headers.get("scd_type") is not None and len(row) > headers["scd_type"] and row[headers["scd_type"]] else "",
        }
        columns.append(col_data)

    wb.close()

    _stm_cache[cache_key] = {"time": now, "columns": columns}
    return columns


def _detect_headers(ws) -> dict | None:
    for row_idx in range(1, 6):
        row = []
        for cell in ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True):
            row = list(cell)
            break

        if not row:
            continue

        found = {"_header_row": row_idx}
        for col_idx, cell_val in enumerate(row):
            if cell_val is None:
                continue
            cell_str = str(cell_val).strip().lower()
            for canonical, variants in KNOWN_HEADERS.items():
                if cell_str in variants and canonical not in found:
                    found[canonical] = col_idx
                    break

        if "target_column" in found:
            return found

    return None
