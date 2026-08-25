from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_xlsx(result: dict, output_path: Path):
    """Generate a two-sheet .xlsx report matching the UI format."""
    wb = Workbook()

    _create_detailed_sheet(wb, result["detailed"])
    _create_summary_sheet(wb, result["summary"])

    wb.save(str(output_path))


def _create_detailed_sheet(wb: Workbook, detailed: list):
    ws = wb.active
    ws.title = "Detailed Comparison"

    headers = [
        "STM Column", "DBT Column", "Column Compare",
        "STM DataType", "DBT DataType", "DataType Compare",
        "STM SCD", "DBT SCD", "SCD Compare",
        "STM Logic", "DBT Logic", "Logic Compare",
        "Suggestion"
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    match_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    mismatch_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_idx, row_data in enumerate(detailed, 2):
        values = [
            row_data.get("stm_column", ""),
            row_data.get("dbt_column", ""),
            row_data.get("column_name_compare", ""),
            row_data.get("stm_datatype", ""),
            row_data.get("dbt_datatype", ""),
            row_data.get("datatype_comparison", ""),
            row_data.get("stm_scd_type", ""),
            row_data.get("dbt_scd_type", ""),
            row_data.get("scd_comparison", ""),
            row_data.get("stm_logic", ""),
            row_data.get("dbt_logic", ""),
            row_data.get("logic_comparison", ""),
            row_data.get("suggestion", ""),
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Color code comparison columns
        for compare_col in [3, 6, 9, 12]:
            cell = ws.cell(row=row_idx, column=compare_col)
            if cell.value and "MATCH" in str(cell.value) and "MISMATCH" not in str(cell.value):
                cell.fill = match_fill
            elif cell.value and "MISMATCH" in str(cell.value):
                cell.fill = mismatch_fill

    col_widths = [20, 20, 28, 15, 15, 16, 10, 10, 14, 40, 40, 14, 20]
    for i, width in enumerate(col_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(64 + (i - 1) % 26 + 1)
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"


def _create_summary_sheet(wb: Workbook, summary: dict):
    ws = wb.create_sheet("Summary")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    bold_font = Font(bold=True)

    ws.cell(row=1, column=1, value="Metric").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=2, value="Value").font = header_font
    ws.cell(row=1, column=2).fill = header_fill

    metrics = [
        ("Total STM Columns", summary.get("total_stm_columns", 0)),
        ("Total DBT Columns", summary.get("total_dbt_columns", 0)),
        ("Matched Columns", summary.get("matched_columns", 0)),
        ("Mismatched Columns", summary.get("datatype_mismatches", 0)),
        ("Extra in DBT", summary.get("extra_in_dbt", 0)),
        ("Match Rate", summary.get("match_rate", "0%")),
    ]

    for row_idx, (metric, value) in enumerate(metrics, 2):
        cell_a = ws.cell(row=row_idx, column=1, value=metric)
        cell_b = ws.cell(row=row_idx, column=2, value=value)
        cell_a.font = bold_font

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
